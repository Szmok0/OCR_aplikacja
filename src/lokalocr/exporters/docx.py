"""DOCX export with python-docx (§5): a clean, editable document, not a page replica."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from lokalocr.exporters.common import paragraph_lines, write_atomic
from lokalocr.model import Document

__all__ = ["export_docx", "to_docx_bytes"]


def to_docx_bytes(document: Document) -> bytes:
    out = DocxDocument()
    normal = out.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for lines in paragraph_lines(document):
        paragraph = out.add_paragraph()
        for i, line in enumerate(lines):
            run = paragraph.add_run(line)
            if i < len(lines) - 1:
                run.add_break(WD_BREAK.LINE)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def export_docx(document: Document, path: Path, *, overwrite: bool = False) -> Path:
    return write_atomic(path, to_docx_bytes(document), overwrite=overwrite)
