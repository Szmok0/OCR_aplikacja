"""Main window tests (stage 3), run offscreen with a controllable fake engine."""

from __future__ import annotations

import io
import shutil
import threading
import time
from pathlib import Path

import pytest

pytest.importorskip("PySide6")

from PIL import Image
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from pytestqt.qtbot import QtBot

from lokalocr import tempdir
from lokalocr.errors import CancelledError, ErrorCode
from lokalocr.model import BBox, Block, BlockType, Document, Line, Page
from lokalocr.ui import strings_pl as S
from lokalocr.ui import theme
from lokalocr.ui.app import create_app
from lokalocr.ui.main_window import MainWindow
from lokalocr.ui.queue_model import Status


class FakeEngine:
    """Returns fixed text; can be slowed down or blocked until cancelled."""

    name = "fake"

    def __init__(self, delay: float = 0.0, block: bool = False) -> None:
        self.delay = delay
        self.block = block
        self.calls = 0

    def recognize(
        self,
        image: Image.Image,
        lang: str,
        *,
        dpi: int,
        page_number: int = 1,
        cancel: threading.Event | None = None,
    ) -> Page:
        self.calls += 1
        deadline = time.monotonic() + (60 if self.block else self.delay)
        while time.monotonic() < deadline:
            if cancel is not None and cancel.wait(0.01):
                raise CancelledError("cancelled")
        lines = [
            Line("Zażółć gęślą jaźń — księgowa", BBox(10, 10, 280, 20), 0.97),
            Line(f"Strona {page_number}", BBox(10, 40, 80, 20), 0.91),
        ]
        return Page(
            page_number,
            image.width,
            image.height,
            dpi,
            [
                Block(
                    BlockType.PARAGRAPH,
                    "\n".join(ln.text for ln in lines),
                    lines,
                    BBox(10, 10, 280, 50),
                    0.94,
                )
            ],
        )

    def searchable_pdf(
        self,
        image: Image.Image,
        lang: str,
        *,
        dpi: int,
        cancel: threading.Event | None = None,
        jpeg_quality: int | None = 80,
    ) -> bytes:
        self.recognize(image, lang, dpi=dpi, cancel=cancel)
        buf = io.BytesIO()
        image.convert("RGB").save(buf, "PDF", resolution=float(dpi))
        return buf.getvalue()

    def available_languages(self) -> frozenset[str]:
        return frozenset({"pol", "eng"})

    def version(self) -> str:
        return "fake"


@pytest.fixture
def files(tmp_path: Path, fixtures_dir: Path) -> Path:
    for name in ("cv.png", "contract.pdf", "invoice.png"):
        shutil.copy(fixtures_dir / name, tmp_path / name)
    return tmp_path


def make_window(qtbot: QtBot, engine: FakeEngine) -> MainWindow:
    create_app(["test"])
    window = MainWindow(lambda: engine)
    qtbot.addWidget(window)
    window.show()
    return window


def wait_idle(qtbot: QtBot, window: MainWindow, timeout: int = 15000) -> None:
    qtbot.waitUntil(
        lambda: (
            not window.controller.busy
            and all(
                i.status in (Status.DONE, Status.ERROR) or i.paused for i in window.queue.items()
            )
        ),
        timeout=timeout,
    )


def test_starts_without_dialogs(qtbot: QtBot) -> None:
    window = make_window(qtbot, FakeEngine())
    assert QApplication.activeModalWidget() is None
    assert window.offline_badge.label.text() == S.OFFLINE_BADGE
    assert window.queue_title.text() == "Kolejka plików (0)"
    assert not window.save_main_btn.isEnabled()
    assert window.result.empty_title.text() == S.EMPTY_RESULT_TITLE


def test_queue_processes_files_sequentially(qtbot: QtBot, files: Path) -> None:
    engine = FakeEngine(delay=0.05)
    window = make_window(qtbot, engine)
    window.add_paths([files])
    assert window.queue_title.text() == "Kolejka plików (3)"
    wait_idle(qtbot, window)
    items = window.queue.items()
    assert [i.status for i in items] == [Status.DONE] * 3
    pdf = next(i for i in items if i.name == "contract.pdf")
    assert pdf.page_count == 2
    qtbot.waitUntil(lambda: all(i.thumbnail is not None for i in items), timeout=5000)

    window.select_item(pdf.id)
    assert "Strona 2" in window.result.visible_text()
    status = window.result.status.text()
    assert "Stron: 2" in status and "Pewność OCR: 94%" in status
    assert "Język OCR: polski + angielski" in status
    assert window.save_main_btn.isEnabled()


def test_layout_switch_and_search(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    window.result.set_layout("simplified")
    assert window.result.layout_mode == "simplified"
    window.result.search.setText("gęślą")
    assert window.result.status.text() == "Wynik 1 z 1"
    assert window.result.text.textCursor().selectedText() == "gęślą"
    window.result.search.setText("nie ma")
    assert window.result.status.text() == S.SEARCH_NO_HITS


def test_copy_text(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    window.result.copy_text()
    assert "Zażółć gęślą jaźń" in QApplication.clipboard().text()


def test_ui_stays_responsive_during_processing(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine(delay=0.4))
    window.add_paths([files / "contract.pdf"])
    started = time.monotonic()
    ticks = 0
    while time.monotonic() - started < 0.6:
        QApplication.processEvents()
        ticks += 1
        time.sleep(0.01)
    assert window.controller.busy
    assert ticks > 30  # the event loop keeps running while OCR works
    item = window.queue.items()[0]
    assert item.status is Status.PROCESSING
    assert window.result.stack.currentIndex() == window.result.PAGE_PROGRESS
    wait_idle(qtbot, window)


def test_cancel_keeps_item_waiting(qtbot: QtBot, files: Path) -> None:
    engine = FakeEngine(block=True)
    window = make_window(qtbot, engine)
    window.add_paths([files / "cv.png"])
    qtbot.waitUntil(lambda: window.controller.busy, timeout=5000)
    item = window.queue.items()[0]
    window.cancel_item(item.id, remove=False)
    qtbot.waitUntil(lambda: item.paused, timeout=5000)
    assert item.status is Status.WAITING
    engine.block = False
    window.reprocess(item.id)
    wait_idle(qtbot, window)
    assert window.queue.items()[0].status is Status.DONE


def test_cancel_and_remove(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine(block=True))
    window.add_paths([files / "cv.png", files / "invoice.png"])
    qtbot.waitUntil(lambda: window.controller.busy, timeout=5000)
    first = window.queue.items()[0]
    window.cancel_item(first.id, remove=True)
    qtbot.waitUntil(lambda: window.queue.get(first.id) is None, timeout=5000)
    # The next file starts automatically.
    qtbot.waitUntil(lambda: window.queue.items()[0].status is Status.PROCESSING, timeout=5000)
    window.controller.cancel()


def test_cancel_dialog_offers_keep_and_remove(
    qtbot: QtBot, files: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = make_window(qtbot, FakeEngine(block=True))
    window.add_paths([files / "cv.png"])
    qtbot.waitUntil(lambda: window.controller.busy, timeout=5000)
    seen: list[str] = []

    def fake_exec(box: QMessageBox) -> int:
        seen.extend(b.text() for b in box.buttons())
        for b in box.buttons():
            if b.text() == S.CANCEL_KEEP:
                b.click()
        return 0

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)
    window.cancel_current()
    assert set(seen) == {S.CANCEL_KEEP, S.CANCEL_REMOVE, S.CANCEL_CONTINUE}
    qtbot.waitUntil(lambda: window.queue.items()[0].paused, timeout=5000)


def test_corrupt_and_unsupported_files(qtbot: QtBot, tmp_path: Path) -> None:
    (tmp_path / "zepsuty.pdf").write_bytes(b"%PDF-1.7 broken")
    (tmp_path / "notatka.png").write_bytes(b"to nie jest obraz")
    window = make_window(qtbot, FakeEngine())
    window.add_paths([tmp_path / "zepsuty.pdf", tmp_path / "notatka.png"])
    assert len(window.queue) == 1  # unsupported content is rejected up front
    assert window.toast.isVisible()
    wait_idle(qtbot, window)
    item = window.queue.items()[0]
    assert item.status is Status.ERROR and item.error_code is ErrorCode.CORRUPT_FILE
    window.select_item(item.id)
    assert window.result.error_text.text() == S.error_message(ErrorCode.CORRUPT_FILE)


def test_duplicates_are_ignored(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    window.add_paths([files / "cv.png"])
    assert len(window.queue) == 1


def test_save_txt(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    out_dir = tmp_path_factory.mktemp("out")
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    suggested: list[str] = []

    def fake_dialog(*args: object, **_k: object) -> tuple[str, str]:
        suggested.append(str(args[2]))
        return str(out_dir / "wynik"), ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", fake_dialog)
    window.save_txt()
    assert suggested[0].endswith("cv_ocr.txt")
    saved = out_dir / "wynik.txt"
    assert "Zażółć gęślą jaźń" in saved.read_text(encoding="utf-8")
    assert window.toast.isVisible()


def test_close_leaves_no_temp_files(qtbot: QtBot, files: Path, tmp_path: Path) -> None:
    root = tmp_path / "lokalocr-temp"
    tempdir.set_root(root)
    try:
        window = make_window(qtbot, FakeEngine())
        window.add_paths([files / "cv.png"])
        wait_idle(qtbot, window)
        (tempdir.app_temp_dir() / "spill.bin").write_bytes(b"x")
        window.close()
        assert not root.exists() or not any(root.iterdir())
        assert len(window.queue) == 0
    finally:
        tempdir.set_root(None)


def test_clear_queue(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files])
    wait_idle(qtbot, window)
    window.clear_queue()
    assert len(window.queue) == 0
    assert window.queue_title.text() == "Kolejka plików (0)"
    assert window.result.stack.currentIndex() == window.result.PAGE_EMPTY


@pytest.mark.parametrize("name", ["light", "dark"])
def test_themes_apply(qtbot: QtBot, name: theme.ThemeName) -> None:
    window = make_window(qtbot, FakeEngine())
    app = QApplication.instance()
    assert isinstance(app, QApplication)
    tokens = theme.apply(app, name)
    assert window.palette().window().color().name().upper() == tokens.bg.upper()
    theme.apply(app, "light")


# --- stage 4: tabs, preview, export ------------------------------------------------


def test_preview_shows_pages_and_navigates(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "contract.pdf"])
    wait_idle(qtbot, window)
    preview = window.preview
    qtbot.waitUntil(lambda: preview._pixmap_item is not None, timeout=5000)
    assert preview.page_total.text() == "/ 2"
    assert not preview.prev_btn.isEnabled() and preview.next_btn.isEnabled()
    preview.next_page()
    assert preview.index == 1
    qtbot.waitUntil(lambda: (files / "contract.pdf", 1) in preview._cache, timeout=5000)
    zoom = preview.view.zoom
    preview.view.zoom_in()
    assert preview.view.zoom > zoom
    assert preview.zoom_label.text() == f"{round(preview.view.zoom * 100)}%"
    preview.rotate(90)
    assert preview.view.rotation == 90


def test_tabs_switch_workspace(qtbot: QtBot, files: Path) -> None:
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    assert not window.tabs.isTabEnabled(2)  # images: full mode only
    assert not window.structure_btn.isEnabled()
    window.tabs.setCurrentIndex(1)
    assert window.result.isHidden()
    qtbot.waitUntil(lambda: bool(window.preview._overlay_items), timeout=5000)
    window.tabs.setCurrentIndex(0)
    assert not window.result.isHidden()
    assert not window.preview._overlay_items


def _save_to(monkeypatch: pytest.MonkeyPatch, target: Path) -> list[str]:
    filters: list[str] = []

    def fake_dialog(*args: object, **_k: object) -> tuple[str, str]:
        filters.append(str(args[3]))
        return str(target), ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", fake_dialog)
    return filters


@pytest.mark.parametrize("fmt", ["docx", "md", "html", "txt"])
def test_save_formats(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
    fmt: str,
) -> None:
    out = tmp_path_factory.mktemp("out") / "wynik"
    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    _save_to(monkeypatch, out)
    window.save(fmt)
    saved = out.with_name(f"wynik.{fmt}")
    assert saved.exists() and saved.stat().st_size > 0
    if fmt != "docx":
        assert "Zażółć gęślą jaźń" in saved.read_text(encoding="utf-8")


def test_save_pdf_in_background(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    out = tmp_path_factory.mktemp("out") / "umowa_ocr.pdf"
    engine = FakeEngine(delay=0.2)
    window = make_window(qtbot, engine)
    window.add_paths([files / "contract.pdf"])
    wait_idle(qtbot, window)
    _save_to(monkeypatch, out)
    window.save("pdf")
    assert window.exporter.busy
    assert window.result.progress_title.text() == S.EXPORTING_TITLE
    qtbot.waitUntil(lambda: not window.exporter.busy, timeout=10000)
    qtbot.waitUntil(lambda: out.exists(), timeout=5000)
    assert out.read_bytes().startswith(b"%PDF")
    assert window.result.stack.currentIndex() == window.result.PAGE_RESULT


def test_cancel_pdf_export(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    out = tmp_path_factory.mktemp("out") / "x.pdf"
    engine = FakeEngine()
    window = make_window(qtbot, engine)
    window.add_paths([files / "contract.pdf"])
    wait_idle(qtbot, window)
    engine.block = True
    _save_to(monkeypatch, out)
    window.save("pdf")
    window.cancel_current()  # Esc while exporting cancels the export
    qtbot.waitUntil(lambda: not window.exporter.busy, timeout=5000)
    assert not out.exists()


# --- stage 5: settings, help, clipboard, themes ---------------------------------------


def test_settings_dialog_roundtrip(qtbot: QtBot, tmp_path: Path) -> None:
    from lokalocr.config import Settings
    from lokalocr.ui.settings_dialog import CATEGORIES, SettingsDialog

    original = Settings(
        default_lang="pol",
        pdf_dpi=200,
        theme="dark",
        output_dir=str(tmp_path),
        clipboard_clear_minutes=5,
    )
    dialog = SettingsDialog(
        original, engine_info=lambda: ("tesseract 5.3.4", frozenset({"pol", "eng"}))
    )
    qtbot.addWidget(dialog)
    assert dialog.categories.count() == len(CATEGORIES)
    assert dialog.collect() == original.validated()
    assert "tesseract 5.3.4" in dialog.lang_status.text()
    dialog.dpi_combo.setCurrentIndex(dialog.dpi_combo.findData(400))
    dialog.theme_radios["light"].setChecked(True)
    dialog.pattern_edit.setText("skan_{nazwa}")
    dialog.accept()
    assert dialog.result_settings.pdf_dpi == 400
    assert dialog.result_settings.theme == "light"
    assert dialog.result_settings.name_pattern == "skan_{nazwa}"


def test_settings_dialog_rejects_bad_pattern(qtbot: QtBot, monkeypatch: pytest.MonkeyPatch) -> None:
    from lokalocr.config import Settings
    from lokalocr.ui.settings_dialog import SettingsDialog

    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: None)
    dialog = SettingsDialog(Settings())
    qtbot.addWidget(dialog)
    dialog.pattern_edit.setText("bez_nazwy")
    dialog.accept()
    assert dialog.result() != dialog.DialogCode.Accepted
    assert dialog.result_settings == Settings()


def test_about_page_is_local_and_offline(qtbot: QtBot) -> None:
    from PySide6.QtWidgets import QTextBrowser

    from lokalocr.config import Settings
    from lokalocr.ui.settings_dialog import SettingsDialog

    dialog = SettingsDialog(Settings())
    qtbot.addWidget(dialog)
    browser = dialog.findChild(QTextBrowser)
    assert browser is not None
    assert not browser.openLinks() and not browser.openExternalLinks()
    assert "Lucide" in browser.toPlainText()


def test_help_window_shows_local_guide(qtbot: QtBot) -> None:
    window = make_window(qtbot, FakeEngine())
    window.open_help()
    assert window._help is not None and window._help.isVisible()
    browser = window._help.browser
    assert not browser.openLinks() and not browser.openExternalLinks()
    assert "Szybki start" in browser.toPlainText()


def test_apply_settings_updates_window(qtbot: QtBot, tmp_path: Path) -> None:
    from lokalocr.config import Settings

    window = make_window(qtbot, FakeEngine())
    window.apply_settings(
        Settings(
            default_lang="eng",
            default_format="md",
            theme="dark",
            auto_preprocess=False,
            keep_structure=False,
            temp_dir=str(tmp_path / "tmp"),
        )
    )
    assert window.lang_combo.currentData() == "eng"
    assert window.save_main_btn.text() == "Zapisz jako Markdown"
    assert not window.preprocess_toggle.isChecked()
    assert not window.keep_structure_cb.isChecked()
    assert window.options().preprocess.enabled is False
    assert tempdir.app_temp_dir() == tmp_path / "tmp"
    assert theme.tokens() is theme.DARK
    window.apply_settings(Settings())
    assert theme.tokens() is not theme.DARK or theme.resolve("system") is theme.DARK


def test_save_to_output_folder_without_dialog(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from lokalocr.config import Settings

    out = tmp_path_factory.mktemp("out")
    window = make_window(qtbot, FakeEngine())
    window.apply_settings(Settings(output_dir=str(out), name_pattern="ocr_{nazwa}"))
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)

    def no_dialog(*_a: object, **_k: object) -> tuple[str, str]:
        raise AssertionError("save dialog must not open")

    monkeypatch.setattr(QFileDialog, "getSaveFileName", no_dialog)
    window.save("txt")
    assert (out / "ocr_cv.txt").exists()
    # Existing file: asks first; "No" keeps the old file.
    (out / "ocr_cv.txt").write_text("stare", encoding="utf-8")
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.No)
    window.save("txt")
    assert (out / "ocr_cv.txt").read_text(encoding="utf-8") == "stare"


def test_clipboard_cleared_after_timeout(qtbot: QtBot, files: Path) -> None:
    from lokalocr.config import Settings

    window = make_window(qtbot, FakeEngine())
    window.apply_settings(Settings(clipboard_clear_minutes=1))
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    window.result.copy_text()
    assert "Zażółć" in QApplication.clipboard().text()
    assert window._clipboard_timer.isActive()
    window._clear_clipboard()  # what the timer does after the delay
    assert QApplication.clipboard().text() == ""


def test_clipboard_not_cleared_if_user_copied_something_else(qtbot: QtBot, files: Path) -> None:
    from lokalocr.config import Settings

    window = make_window(qtbot, FakeEngine())
    window.apply_settings(Settings(clipboard_clear_minutes=1))
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    window.result.copy_text()
    QApplication.clipboard().setText("coś innego")
    window._clear_clipboard()
    assert QApplication.clipboard().text() == "coś innego"


def test_ctrl_c_without_selection_copies_all(qtbot: QtBot, files: Path) -> None:
    from PySide6.QtCore import Qt as QtCore_Qt

    window = make_window(qtbot, FakeEngine())
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    QApplication.clipboard().clear()
    window.result.text.setFocus()
    qtbot.keyClick(  # type: ignore[no-untyped-call]
        window.result.text, QtCore_Qt.Key.Key_C, QtCore_Qt.KeyboardModifier.ControlModifier
    )
    assert "Zażółć gęślą jaźń" in QApplication.clipboard().text()


def test_engine_banner(qtbot: QtBot) -> None:
    from lokalocr.errors import EngineMissingError

    def missing() -> FakeEngine:
        raise EngineMissingError("no tesseract")

    create_app(["test"])
    window = MainWindow(missing)
    qtbot.addWidget(window)
    window.show()
    window.check_engine()
    qtbot.waitUntil(lambda: window.banner.isVisible(), timeout=5000)

    ok = make_window(qtbot, FakeEngine())
    ok.check_engine()
    qtbot.wait(200)
    assert not ok.banner.isVisible()


# --- stage 6: full mode ----------------------------------------------------------------


class FakeStructureEngine:
    """Returns a structured document like Docling would (no models)."""

    name = "fake-structure"

    def __init__(self, workers: int = 1) -> None:
        self.workers = workers

    def is_available(self) -> bool:
        return True

    def analyze(
        self,
        path: Path,
        lang: str,
        *,
        cancel: threading.Event | None = None,
        on_progress: object = None,
    ) -> Document:
        from lokalocr.model import Table, TableCell

        if callable(on_progress):
            on_progress(0, 1)
            on_progress(1, 1)
        png = io.BytesIO()
        Image.new("RGB", (60, 40), (20, 120, 240)).save(png, "PNG")
        blocks = [
            Block(BlockType.HEADING, "Zofia Przykładowa", level=1, bbox=BBox(50, 40, 250, 20)),
            Block(BlockType.HEADING, "Doświadczenie", level=2, bbox=BBox(50, 80, 150, 15)),
            Block(
                BlockType.PARAGRAPH,
                "Księgowa z doświadczeniem.",
                bbox=BBox(50, 100, 400, 30),
                confidence=0.93,
            ),
            Block(BlockType.LIST_ITEM, "Rozliczenia VAT", bbox=BBox(60, 140, 200, 12)),
            Block(BlockType.LIST_ITEM, "Płace", bbox=BBox(60, 155, 200, 12)),
            Block(
                BlockType.TABLE,
                bbox=BBox(50, 200, 400, 60),
                table=Table(
                    2,
                    2,
                    [
                        TableCell(0, 0, "Nazwa"),
                        TableCell(0, 1, "Cena"),
                        TableCell(1, 0, "Segregatory"),
                        TableCell(1, 1, "35,50"),
                    ],
                ),
            ),
            Block(BlockType.FIGURE, bbox=BBox(50, 300, 120, 80), image_ref="figure_1.png"),
            Block(BlockType.CAPTION, "Rys. 1. Logo", bbox=BBox(50, 385, 120, 10)),
        ]
        return Document(
            pages=[Page(1, 595, 842, 72, blocks)],
            lang=lang,
            mode="full",
            images={"figure_1.png": png.getvalue()},
        )


def make_full_window(qtbot: QtBot) -> MainWindow:
    create_app(["test"])
    window = MainWindow(
        lambda: FakeEngine(), full_mode_available=True, structure_factory=FakeStructureEngine
    )
    qtbot.addWidget(window)
    window.show()
    return window


def test_full_mode_document(qtbot: QtBot, files: Path) -> None:
    window = make_full_window(qtbot)
    assert window.full_radio.isEnabled()
    window.full_radio.setChecked(True)
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    item = window.queue.items()[0]
    assert item.status is Status.DONE and item.document is not None
    assert item.document.mode == "full"
    assert window.structure_btn.isEnabled()
    assert window.keep_structure_cb.isEnabled()
    assert window.tabs.isTabEnabled(2) and window.tabs.tabText(2) == "Obrazy (1)"
    window.tabs.setCurrentIndex(2)
    assert window.gallery.count() == 1
    assert window.open_images_btn.isEnabled()


def test_structure_window(qtbot: QtBot, files: Path) -> None:
    window = make_full_window(qtbot)
    window.full_radio.setChecked(True)
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    window.open_structure()
    structure = window._structure
    assert structure is not None and structure.isVisible()
    assert structure.windowTitle() == "Struktura dokumentu – cv.png"
    page = structure.tree.topLevelItem(0)
    assert page is not None and page.text(0) == "Strona 1"
    labels: list[str] = []

    def walk(node: object) -> None:
        from PySide6.QtWidgets import QTreeWidgetItem

        assert isinstance(node, QTreeWidgetItem)
        for i in range(node.childCount()):
            child = node.child(i)
            if child is not None:
                labels.append(child.text(0))
                walk(child)

    walk(page)
    assert "Sekcja: Zofia Przykładowa" in labels
    assert "Sekcja: Doświadczenie" in labels
    assert "Lista (2)" in labels
    assert any(lbl.startswith("Tabela") for lbl in labels)
    assert "Obraz" in labels
    # Select the paragraph: properties are shown and the main preview highlights it.
    from PySide6.QtWidgets import QTreeWidgetItemIterator

    it = QTreeWidgetItemIterator(structure.tree)
    while it.value() is not None:
        node = it.value()
        if node.text(0).startswith("Akapit:"):
            structure.tree.setCurrentItem(node)
            break
        it += 1
    assert structure.type_label.text() == "Akapit"
    assert structure.conf_label.text() == "93%"
    assert window.preview._highlight == BBox(50, 100, 400, 30)


def test_open_images_folder(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from PySide6.QtGui import QDesktopServices

    out = tmp_path_factory.mktemp("out")
    window = make_full_window(qtbot)
    window.full_radio.setChecked(True)
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    opened: list[str] = []
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", lambda *a, **k: str(out))
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda url: opened.append(url.toLocalFile()))
    window.open_images_folder()
    assert (out / "obrazy" / "cv_ocr_figure_1.png").exists()
    assert [Path(p) for p in opened] == [out / "obrazy"]


def test_full_mode_docx_has_structure(
    qtbot: QtBot,
    files: Path,
    tmp_path_factory: pytest.TempPathFactory,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from docx import Document as DocxDocument

    out = tmp_path_factory.mktemp("out") / "cv_ocr.docx"
    window = make_full_window(qtbot)
    window.full_radio.setChecked(True)
    window.add_paths([files / "cv.png"])
    wait_idle(qtbot, window)
    _save_to(monkeypatch, out)
    window.save("docx")
    doc = DocxDocument(str(out))
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert any(p.style is not None and p.style.name == "Heading 1" for p in doc.paragraphs)
    # "Zachowaj strukturę" off: plain paragraphs only.
    window.keep_structure_cb.setChecked(False)
    out.unlink()
    window.save("docx")
    plain = DocxDocument(str(out))
    assert not plain.tables and not plain.inline_shapes
