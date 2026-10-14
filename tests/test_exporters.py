from __future__ import annotations

import io
from pathlib import Path

import pypdfium2 as pdfium
import pytest
from docx import Document as DocxDocument

import make_fixtures as fx
from lokalocr.exporters.common import OutputExistsError
from lokalocr.exporters.docx import export_docx, to_docx_bytes
from lokalocr.exporters.html import CSP, to_html
from lokalocr.exporters.markdown import to_markdown
from lokalocr.model import BBox, Block, BlockType, Document, Line, Page
from measure_cer import cer


def tricky_document() -> Document:
    """Text that would read as markup in Markdown or HTML."""
    blocks = [
        Block(BlockType.PARAGRAPH, "Księgowa z doświadczeniem *gwiazdka* i [nawias]."),
        Block(BlockType.PARAGRAPH, "Płace <i> & kadry <script>alert(1)</script>"),
        Block(BlockType.PARAGRAPH, '<img src="https://example.com/x.png">'),
    ]
    return Document(pages=[Page(1, 100, 100, 300, blocks)])


def basic_document() -> Document:
    lines = [
        Line("Przykładowa Firma Sp. z o.o.", BBox(0, 0, 280, 20)),
        Line("ul. Łąkowa 12, 80-001 Gdańsk", BBox(0, 30, 290, 20)),
    ]
    return Document(
        pages=[
            Page(
                1,
                100,
                100,
                300,
                [
                    Block(
                        BlockType.PARAGRAPH,
                        "\n".join(ln.text for ln in lines),
                        lines,
                        BBox(0, 0, 290, 50),
                    ),
                    Block(BlockType.PARAGRAPH, "Zażółć gęślą jaźń"),
                ],
            )
        ]
    )


def test_docx_is_plain_paragraphs(tmp_path: Path) -> None:
    doc = DocxDocument(str(export_docx(tricky_document(), tmp_path / "a.docx")))
    assert not doc.tables and not doc.inline_shapes
    assert {p.style.name if p.style else "" for p in doc.paragraphs} == {"Normal"}
    assert doc.paragraphs[1].text == "Płace <i> & kadry <script>alert(1)</script>"


def test_docx_basic_keeps_address_lines() -> None:
    doc = DocxDocument(io.BytesIO(to_docx_bytes(basic_document())))
    first = doc.paragraphs[0]
    assert first.text == "Przykładowa Firma Sp. z o.o.\nul. Łąkowa 12, 80-001 Gdańsk"
    assert doc.paragraphs[1].text == "Zażółć gęślą jaźń"


def test_docx_never_overwrites(tmp_path: Path) -> None:
    target = tmp_path / "a.docx"
    target.write_bytes(b"old")
    with pytest.raises(OutputExistsError):
        export_docx(basic_document(), target)
    assert target.read_bytes() == b"old"


def test_markdown_escaping() -> None:
    md = to_markdown(tricky_document())
    assert "\\*gwiazdka\\*" in md and "\\[nawias\\]" in md


def test_markdown_escapes_line_start_markers() -> None:
    doc = Document(pages=[Page(1, 1, 1, 300, [Block(BlockType.PARAGRAPH, "# nie nagłówek")])])
    assert to_markdown(doc).startswith("\\# nie nagłówek")


def test_markdown_basic_line_breaks() -> None:
    assert to_markdown(basic_document()).startswith(
        "Przykładowa Firma Sp. z o.o.  \nul. Łąkowa 12, 80-001 Gdańsk\n\nZażółć"
    )


def test_html_escaping_and_no_external_resources() -> None:
    html = to_html(tricky_document(), title="cv <x>")
    assert '<meta charset="utf-8">' in html and "<title>cv &lt;x&gt;</title>" in html
    assert "<p>Płace &lt;i&gt; &amp; kadry &lt;script&gt;alert(1)&lt;/script&gt;</p>" in html
    assert "<script" not in html and "<img" not in html
    # The browser may not load anything when the file is opened (§6).
    assert f'<meta http-equiv="Content-Security-Policy" content="{CSP}">' in html
    assert CSP.startswith("default-src 'none'")


def test_html_basic_paragraphs() -> None:
    html = to_html(basic_document())
    assert "<p>Przykładowa Firma Sp. z o.o.<br>\nul. Łąkowa 12, 80-001 Gdańsk</p>" in html


@pytest.mark.tesseract
def test_searchable_pdf(tmp_path: Path) -> None:
    from lokalocr.engines.tesseract import TesseractEngine
    from lokalocr.exporters.pdf_text import export_searchable_pdf
    from lokalocr.pipeline import ProcessOptions

    contract = fx.make_contract(150, pages=2)
    source = fx.save(contract, tmp_path / "umowa.pdf")
    events: list[str] = []
    out = export_searchable_pdf(
        source,
        tmp_path / "umowa_ocr.pdf",
        TesseractEngine(),
        ProcessOptions(pdf_dpi=150),
        on_progress=lambda e: events.append(e.stage),
    )
    assert set(events) == {"export"}
    pdf = pdfium.PdfDocument(str(out))
    try:
        assert len(pdf) == 2
        width, height = pdf[0].get_size()
        assert abs(width - 595) < 3 and abs(height - 842) < 3  # A4 kept
        text = "".join(pdf[i].get_textpage().get_text_range() for i in range(2))
        assert cer(contract.truth, text) < 0.05
        assert pdf[1].get_textpage().search("Wykonawca").get_next() is not None
    finally:
        pdf.close()
    # JPEG-embedded pages stay small.
    assert out.stat().st_size < 1_000_000 * 2
